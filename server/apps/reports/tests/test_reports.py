import csv
import io

import pytest
from openpyxl import load_workbook
from rest_framework.test import APIClient

from apps.billing.models import Folio, Payment
from apps.cash.models import Expense, Shift
from apps.stays.models import Reservation

pytestmark = pytest.mark.django_db


def get(api, name, **params):
    res = api.get(f"/api/v1/reports/{name}", params)
    assert res.status_code == 200, res.json()
    return res.json()


def test_index_lists_reports_with_badges(api_as_manager):
    index = {r["name"]: r for r in api_as_manager.get("/api/v1/reports/").json()}
    assert {
        "occupancy",
        "arrivals_departures",
        "current_guests",
        "ending_soon",
        "debts",
        "revenue",
        "expenses",
        "cash_shifts",
        "adjustments",
        "room_status",
    } <= set(index)
    assert index["debts"]["title"] == "الديون" and index["debts"]["badge"] == 2
    assert index["ending_soon"]["badge"] == 2  # 207 and 305 overdue


def test_debts_match_the_brief(api_as_manager):
    report = get(api_as_manager, "debts")
    assert {r["room"]: r["balance"] for r in report["rows"]} == {"203": 1_500_000, "305": 4_200_000}
    tiles = {t["label"]: t for t in report["meta"]["tiles"]}
    assert tiles["الديون المستحقة"]["value"] == 5_700_000  # «57,000 · إقامتان جاريتان»
    assert tiles["الديون المستحقة"]["hint"] == "2 إقامات جارية"
    assert report["meta"]["totals"]["balance"] == 5_700_000
    assert [c["label"] for c in report["columns"]][-2:] == ["عمر الدين", "الحالة / السبب"]


def test_arrivals_and_departures(api_as_manager):
    today = get(api_as_manager, "arrivals_departures")
    moves = sorted((r["move"], r["room"], r["ready"]) for r in today["rows"])
    # 301 arrived today (weekly, ends in 6 days) and is already in; 108/204 end today; 207/305 are overdue.
    assert moves == [
        ("مغادرة", "108", "—"),
        ("مغادرة", "204", "—"),
        ("مغادرة", "207", "—"),
        ("مغادرة", "305", "—"),
        ("وصول", "301", "مسكّن"),
    ]
    tiles = {t["label"]: t["value"] for t in today["meta"]["tiles"]}
    assert tiles == {"وصول اليوم": 1, "مغادرة اليوم": 4, "غرف تحتاج تجهيزًا قبل الوصول": 0}
    tomorrow = get(api_as_manager, "arrivals_departures", when="tomorrow")
    arrival = next(r for r in tomorrow["rows"] if r["move"] == "وصول")
    assert arrival["room"] == "102" and arrival["ready"] == "جاهزة"


def test_occupancy_tonight(api_as_manager):
    report = get(api_as_manager, "occupancy")
    tonight = report["rows"][-1]
    assert (tonight["occupied"], tonight["available"], tonight["maintenance"]) == (18, 29, "410")
    assert tonight["occupancy"] == 62  # 18 ÷ (30 − 1)
    assert "منتصف الليل" in report["meta"]["formula"]


def test_guest_lists(api_as_manager):
    assert len(get(api_as_manager, "current_guests")["rows"]) == 18
    soon = get(api_as_manager, "ending_soon")["rows"]
    assert [r["room"] for r in soon][:2] == ["305", "207"]
    assert soon[0]["state"] == "متجاوزة منذ يومين"


def test_money_reports(api_as_manager):
    revenue = get(api_as_manager, "revenue")
    lines_total = sum(
        f.lines.aggregate(s=__import__("django.db.models", fromlist=["Sum"]).Sum("amount"))["s"] or 0
        for f in Folio.objects.all()
    )
    assert revenue["meta"]["totals"]["revenue"] == lines_total
    assert revenue["meta"]["totals"]["collected"] == sum(p.amount for p in Payment.objects.all())
    expenses = get(api_as_manager, "expenses")
    assert expenses["meta"]["totals"]["amount"] == 1_250_000 and len(expenses["rows"]) == 2
    shifts = get(api_as_manager, "cash_shifts")
    assert len(shifts["rows"]) == Shift.objects.count() == 2
    assert "البنكك والتحويل لا يدخلان الدرج" in shifts["meta"]["formula"]
    assert len(get(api_as_manager, "room_status")["rows"]) == 30


def test_adjustments_collects_discounts_and_reversals(api_as_manager):
    folio = Reservation.objects.get(room__number="203", status="checked_in").folio
    res = api_as_manager.post(
        f"/api/v1/folios/{folio.pk}/lines",
        {"kind": "discount", "amount": 500_000, "reason": "نزيل دائم"},
        format="json",
    )
    assert res.status_code == 201, res.json()
    expense = Expense.objects.first()
    api_as_manager.post(f"/api/v1/expenses/{expense.pk}/reverse", {"reason": "خطأ"}, format="json")
    report = get(api_as_manager, "adjustments")
    assert {r["type"] for r in report["rows"]} == {"خصم", "عكس مصروف"}
    assert report["meta"]["tiles"][0]["value"] == 500_000


def test_bad_requests(api_as_manager):
    assert api_as_manager.get("/api/v1/reports/nope").status_code == 404
    assert api_as_manager.get("/api/v1/reports/debts", {"date_from": "26-09-2026"}).json()["code"] == "validation_error"
    assert (
        api_as_manager.get("/api/v1/reports/debts/export", {"format": "docx"}).status_code == 400
    )  # pdf is a format since 1.1.14


def test_xlsx_export_is_rtl_with_numeric_money(api_as_manager):
    res = api_as_manager.get("/api/v1/reports/debts/export", {"format": "xlsx"})
    assert res.status_code == 200
    assert res["Content-Disposition"].startswith('attachment; filename="skytowers-debts-')
    ws = load_workbook(io.BytesIO(res.content)).active
    assert ws.sheet_view.rightToLeft is True
    header = [c.value for c in ws[4]]
    assert header[0] == "النزيل" and "المتبقي" in header
    balances = {row[1]: row[5] for row in ws.iter_rows(min_row=5, values_only=True) if row[1] in ("203", "305")}
    assert balances == {"203": 15000, "305": 42000}  # pounds, as numbers
    assert ws.freeze_panes == "A5"


def test_csv_export_has_bom_and_crlf(api_as_manager):
    res = api_as_manager.get("/api/v1/reports/debts/export", {"format": "csv"})
    raw = res.content
    assert raw.startswith("﻿".encode()) and b"\r\n" in raw
    rows = list(csv.reader(io.StringIO(raw.decode("utf-8-sig"))))
    assert rows[0][0] == "النزيل"
    assert {r[1]: r[5] for r in rows[1:]} == {"203": "15000", "305": "42000"}


class TestDocuments:
    def test_invoice(self, api_as_manager):
        folio = Reservation.objects.get(room__number="203", status="checked_in").folio
        doc = api_as_manager.get(f"/api/v1/folios/{folio.pk}/invoice").json()
        assert doc["invoice"] == folio.invoice_label
        assert doc["guest"]["name"] == "محمد عثمان الطيب" and doc["guest"]["id_number"] == "211-8842-1023-7"
        assert doc["totals"]["balance"] == 1_500_000
        # Artboard 7.1: line items net of reversals, the discount in the totals block, the payment list.
        assert sum(i["total"] for i in doc["items"]) == doc["totals"]["charges"]
        assert doc["totals"]["charges"] - doc["totals"]["discount"] == doc["totals"]["total"]
        assert all(
            i["quantity"] >= 1 and i["unit_price"] * i["quantity"] <= i["total"] + i["quantity"] for i in doc["items"]
        )
        assert sum(p["amount"] for p in doc["payments"]) == doc["totals"]["paid"]
        assert doc["stay"]["state"].startswith("جارية — تنتهي بنهاية يوم")
        assert doc["hotel"]["name_ar"] == "فندق سكاي تاورز"

    def test_payment_receipt(self, api_as_manager):
        payment = Payment.objects.first()
        doc = api_as_manager.get(f"/api/v1/payments/{payment.pk}/receipt").json()
        assert doc["receipt"] == payment.receipt_label and doc["amount"] == payment.amount
        assert doc["amount_in_words"].endswith("جنيه") and doc["balance"] >= 0
        assert doc["nights"] >= 1 and doc["check_in_date"] <= doc["last_night"]  # design gap #11: the stay period

    def test_expense_receipt_and_shift_statement(self, api_as_manager):
        expense = Expense.objects.order_by("number").first()
        doc = api_as_manager.get(f"/api/v1/expenses/{expense.pk}/receipt").json()
        assert doc["number"] == "EXP-000001" and doc["method"] == "نقدي — من الدرج"
        shift = Shift.objects.filter(closed_at__isnull=True).get()
        statement = api_as_manager.get(f"/api/v1/shifts/{shift.pk}/statement").json()
        assert statement["tiles"]["expected"] == 5_000_000 - 1_250_000
        assert [m["kind"] for m in statement["movements"]].count("out") == 2


def test_invoice_masks_the_id_number_for_reception(seeded):
    """Spec §5: the guest id number is visible to manager/owner only — the printed invoice included."""
    from apps.accounts.services import login_with_password
    from apps.core.seed import demo_data

    client = APIClient()
    client.credentials(HTTP_AUTHORIZATION=f"Token {login_with_password('ahmed.ali', demo_data.DEMO_PASSWORD).token}")
    folio = Reservation.objects.get(room__number="203", status="checked_in").folio
    doc = client.get(f"/api/v1/folios/{folio.pk}/invoice").json()
    assert doc["guest"]["id_number"] == "••••23-7"


def test_report_filters_narrow_rows_and_are_echoed_in_meta(api_as_manager):
    from apps.rooms.models import RoomType

    double = RoomType.objects.get(name="مزدوجة")
    everyone = get(api_as_manager, "current_guests")["rows"]
    doubles = get(api_as_manager, "current_guests", room_type=str(double.pk))
    assert 0 < len(doubles["rows"]) < len(everyone)
    assert doubles["meta"]["filters"] == [{"key": "room_type", "label": "نوع الغرفة", "value": "مزدوجة"}]
    assert get(api_as_manager, "current_guests")["meta"]["filters"] == []

    cash_only = get(api_as_manager, "revenue", method="cash")
    assert cash_only["meta"]["totals"]["bankak"] == 0 and cash_only["meta"]["totals"]["transfer"] == 0
    assert cash_only["meta"]["totals"]["collected"] == cash_only["meta"]["totals"]["cash"]
    assert cash_only["meta"]["filters"][0]["value"] == "نقدي"

    supplies = get(api_as_manager, "expenses", expense_category="supplies")["rows"]
    assert supplies and all(r["category"] == "مستلزمات" for r in supplies)

    for name, bad in (
        ("current_guests", {"room_type": "not-a-uuid"}),
        ("expenses", {"method": "gold"}),
        ("expenses", {"expense_category": "nope"}),
    ):
        res = api_as_manager.get(f"/api/v1/reports/{name}", bad)
        assert res.status_code == 400 and res.json()["code"] == "validation_error", bad
